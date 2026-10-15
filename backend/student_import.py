"""
O'quvchilar ro'yxatini Excel (.xlsx), CSV yoki joylangan matndan o'qish.

Maqsad — o'qituvchi tayyor ro'yxatini qayta yozib o'tirmasin. Jadval
qanday tuzilganidan qat'i nazar ismlar ajratib olinadi:

  • sarlavha qatori ("№", "F.I.Sh", "Ism", "Familiya") tashlab ketiladi;
  • raqamlar (tartib raqami), sanalar va telefon raqamlari e'tiborsiz;
  • "Familiya | Ism | Otasining ismi" kabi alohida ustunlar bitta ismga
    birlashtiriladi;
  • sarlavhada ism ustunlari ko'rsatilgan bo'lsa ("F.I.Sh", "Ism", "Familiya"),
    faqat o'shalar olinadi — "Sinf", "Manzil" kabi boshqa ustunlar ismga qo'shilmaydi;
  • takrorlar olib tashlanadi.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime

MAX_STUDENTS = 500
MAX_NAME_PARTS = 3

HEADER_WORDS = (
    "f.i.sh", "fish", "f.i.o", "fio", "ism", "familiya", "sharif", "o'quvchi", "oquvchi", "o‘quvchi",
    "name", "student", "ф.и.о", "фио", "имя", "фамилия", "ученик", "№",
)
_LETTER = re.compile(r"[^\W\d_]", re.UNICODE)
_NOT_NAME = re.compile(r"^[\d\s.,:;/\\\-+()№#]+$")  # faqat raqam/belgi: "1.", "14.05.2010", "+998 90 ..."


def _is_name_cell(value) -> bool:
    if value is None or isinstance(value, (int, float, date, datetime)):
        return False
    text = str(value).strip()
    if len(text) < 2 or _NOT_NAME.match(text):
        return False
    return bool(_LETTER.search(text))


def _is_header(row: list) -> bool:
    cells = [str(cell).strip().lower() for cell in row if cell not in (None, "")]
    return any(word in cell for cell in cells for word in HEADER_WORDS) and len(cells) <= 8


# Sarlavhadagi ism ustunlari (F.I.Sh., Ism, Familiya, Otasining ismi ...)
NAME_COLUMN_WORDS = (
    "f.i.sh", "fish", "f.i.o", "fio", "ism", "familiya", "sharif", "o'quvchi", "oquvchi", "o‘quvchi",
    "name", "student", "ф.и.о", "фио", "имя", "фамилия", "отчество", "ученик",
)


def _skipped_columns(header: list) -> set[int] | None:
    """
    Sarlavhada nomi bor, lekin ismga aloqasi yo'q ustunlar ("Sinf", "Manzil", "Telefon") —
    ular ismga qo'shilmaydi. Sarlavhasiz ustunlar (masalan "F.I.Sh" ostida familiya va ism
    alohida katakda) olinaveradi. Sarlavhada ism ustuni bo'lmasa — hech narsa tashlanmaydi.
    """

    labelled = {index: str(cell).strip().lower() for index, cell in enumerate(header) if cell not in (None, "")}
    is_name = {index: any(word in text for word in NAME_COLUMN_WORDS) for index, text in labelled.items()}
    if not any(is_name.values()):
        return None
    return {index for index, name_column in is_name.items() if not name_column and labelled[index] not in ("№", "#")}


def names_from_rows(rows: list[list]) -> list[str]:
    names: list[str] = []
    seen: set[str] = set()
    skipped: set[int] = set()

    for index, row in enumerate(rows):
        if not row:
            continue
        # Sarlavha odatda birinchi qatorlarda bo'ladi
        if index < 3 and _is_header(row):
            skipped = _skipped_columns(row) or set()
            continue
        cells = [cell for column, cell in enumerate(row) if column not in skipped]
        parts = [" ".join(str(cell).split()) for cell in cells if _is_name_cell(cell)][:MAX_NAME_PARTS]
        if not parts:
            continue
        name = " ".join(parts)[:200]
        key = name.lower()
        if len(name) >= 2 and key not in seen:
            seen.add(key)
            names.append(name)
        if len(names) >= MAX_STUDENTS:
            break
    return names


def rows_from_text(text: str) -> list[list[str]]:
    """Joylangan matn: har qator — bitta o'quvchi. Excel'dan nusxalanganda ustunlar TAB bilan ajraladi."""
    rows = []
    for line in (text or "").splitlines():
        if "\t" in line:
            rows.append(line.split("\t"))
        elif ";" in line:
            rows.append(line.split(";"))
        else:
            # "1. Karimov Islom" yoki "1) Karimov Islom" — boshidagi raqamni olib tashlaymiz
            rows.append([re.sub(r"^\s*\d+\s*[.)\-]?\s*", "", line)])
    return rows


def rows_from_csv(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    except csv.Error:
        return rows_from_text(text)


def rows_from_xlsx(data: bytes) -> list[list]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True, max_row=MAX_STUDENTS + 20)]
    finally:
        workbook.close()


def parse_student_file(file_name: str, data: bytes) -> list[str]:
    """Fayl turiga qarab ismlarni ajratadi. Noma'lum/buzuq fayl — ValueError."""

    name = (file_name or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        raise ValueError("Eski .xls format qo'llanmaydi — Excel'da «Saqlash» → .xlsx qilib qayta saqlang")
    try:
        if name.endswith(".xlsx") or data[:2] == b"PK":
            rows = rows_from_xlsx(data)
        elif name.endswith((".csv", ".txt")):
            rows = rows_from_csv(data)
        else:
            raise ValueError("Faqat .xlsx, .csv yoki .txt fayl yuklang")
    except ValueError:
        raise
    except Exception as error:  # buzuq fayl
        raise ValueError("Faylni o'qib bo'lmadi — u buzilgan yoki boshqa formatda") from error
    return names_from_rows(rows)
